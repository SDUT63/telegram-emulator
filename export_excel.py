#!/usr/bin/env python3
"""
Выгрузка обращений в Excel с напоминаниями о контрольных звонках.

Запуск:
    python export_excel.py

Получится файл «Обращения СДУТ <дата>.xlsx» с двумя листами:

    Обращения — все анкеты целиком, статусы, кто ведёт, заметки
    Звонки    — контрольные звонки через 7 и 30 дней, с подсветкой
                просроченных

Скрипт только читает данные и ничего не меняет: его можно запускать
в любой момент, не останавливая ни бота, ни CRM.
"""

from __future__ import annotations

import os
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import crm_store as store
from chatbot_survey import Survey
from survey_questions import QUESTIONS

# Фирменные цвета: тот же петроль и песочный, что на страницах
PETROL = "1F5257"
SAND = "C09A6B"
CREAM = "F6F1E8"
RED = "A93E32"
RED_SOFT = "F8E7E3"
AMBER_SOFT = "F6EBD8"
GREEN_SOFT = "DFEAE6"

HEAD_FILL = PatternFill("solid", fgColor=PETROL)
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="E4DACA")
BORDER = Border(bottom=THIN)


def parse(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def call_state(due: datetime, done: dict | None, today: datetime) -> tuple[str, str]:
    """Состояние звонка и цвет строки."""
    if done:
        return "Сделан " + (parse(done.get("at")) or today).strftime("%d.%m.%Y"), GREEN_SOFT
    days = (due.date() - today.date()).days
    if days < 0:
        return f"Просрочен на {-days} дн.", RED_SOFT
    if days == 0:
        return "Сегодня", AMBER_SOFT
    if days <= 2:
        return f"Через {days} дн.", AMBER_SOFT
    return f"Через {days} дн.", ""


def style_header(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"


def build(path: str | None = None) -> str:
    survey = Survey()          # только чтение
    cases = store.all_cases()
    delivered = store.delivery_state()
    today = datetime.now()

    book = Workbook()

    # ------------------------------------------------ лист «Звонки»
    # Первым, потому что с него начинают день
    calls_sheet = book.active
    calls_sheet.title = "Звонки"
    calls_sheet.append(
        ["Когда звонить", "Состояние", "Имя", "Телефон", "Через сколько дней",
         "Статус обращения", "Кто ведёт", "Признаки"]
    )

    rows: list[tuple] = []
    for user_id, person in survey.state.items():
        answers = person.get("answers", {})
        operator = cases.get(user_id, {})
        base = parse(person.get("finished")) or parse(person.get("started"))
        if not base:
            continue
        for which, days in store.CALL_STAGES.items():
            due = base + timedelta(days=days)
            done = (operator.get("calls") or {}).get(which)
            label, colour = call_state(due, done, today)
            rows.append(
                (
                    due,
                    label,
                    colour,
                    answers.get("name", "без имени"),
                    answers.get("phone", ""),
                    int(which),
                    operator.get("status", store.STATUSES[0]),
                    operator.get("assigned", ""),
                    "; ".join(person.get("alerts", [])),
                )
            )

    # Сначала просроченные и ближайшие
    rows.sort(key=lambda r: (bool(r[2] == GREEN_SOFT), r[0]))

    for due, label, colour, name, phone, days, status, who, alerts in rows:
        calls_sheet.append(
            [due.strftime("%d.%m.%Y"), label, name, phone, days, status, who, alerts]
        )
        if colour:
            for cell in calls_sheet[calls_sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=colour)
        for cell in calls_sheet[calls_sheet.max_row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    style_header(calls_sheet, [15, 20, 22, 16, 12, 16, 16, 40])
    if calls_sheet.max_row > 1:
        calls_sheet.auto_filter.ref = f"A1:H{calls_sheet.max_row}"

    # ------------------------------------------------ лист «Обращения»
    sheet = book.create_sheet("Обращения")
    header = [
        "Обратился", "Заполнено", "Статус", "Кто ведёт", "Требует внимания",
        "Звонок 7 дней", "Звонок 30 дней", "Заметки", "Написано в чат",
    ] + [q["text"].splitlines()[0] for q in QUESTIONS]
    sheet.append(header)

    for user_id, person in survey.state.items():
        answers = person.get("answers", {})
        operator = cases.get(user_id, {})
        calls = operator.get("calls") or {}
        notes = "\n".join(
            f"{n.get('at','')[:16].replace('T',' ')} {n.get('who','')}: {n.get('text','')}"
            for n in operator.get("notes", [])
        )
        sent = "\n".join(
            f"{m.get('who','')}: {m.get('text','')}"
            f" [{'доставлено' if delivered.get(m.get('id'), {}).get('delivered') else 'в очереди'}]"
            for m in operator.get("sent", [])
        )
        alerts = "; ".join(person.get("alerts", []))
        sheet.append(
            [
                (person.get("started") or "")[:16].replace("T", " "),
                "да" if person.get("finished") else "нет",
                operator.get("status", store.STATUSES[0]),
                operator.get("assigned", ""),
                alerts,
                "сделан" if calls.get("7") else "",
                "сделан" if calls.get("30") else "",
                notes,
                sent,
            ]
            + [answers.get(q["id"], "") for q in QUESTIONS]
        )
        row = sheet[sheet.max_row]
        if alerts:
            row[4].font = Font(color=RED, bold=True)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER

    style_header(sheet, [17, 11, 12, 14, 34, 13, 14, 40, 34] + [26] * len(QUESTIONS))
    if sheet.max_row > 1:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{sheet.max_row}"

    path = path or os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        f"Обращения СДУТ {today.strftime('%Y-%m-%d')}.xlsx",
    )
    book.save(path)
    return path


def _cli() -> int:
    survey = Survey()
    if not survey.state:
        print()
        print("  Обращений пока нет — выгружать нечего.")
        print()
        return 1

    path = build()
    total = len(survey.state)
    today = datetime.now()

    overdue = 0
    for user_id, person in survey.state.items():
        base = parse(person.get("finished")) or parse(person.get("started"))
        if not base:
            continue
        calls = (store.all_cases().get(user_id, {}).get("calls") or {})
        for which, days in store.CALL_STAGES.items():
            if calls.get(which):
                continue
            if (base + timedelta(days=days)).date() < today.date():
                overdue += 1

    print()
    print("=" * 62)
    print(f"  Готово: {os.path.basename(path)}")
    print()
    print(f"  Обращений в файле: {total}")
    if overdue:
        print(f"  Просроченных контрольных звонков: {overdue}")
        print("  Они на первом листе, подсвечены красным.")
    else:
        print("  Просроченных контрольных звонков нет.")
    print()
    print("  Открыть: ii \"" + os.path.basename(path) + "\"")
    print("=" * 62)
    print()
    return 0


if __name__ == "__main__":
    sys.exit(_cli())
